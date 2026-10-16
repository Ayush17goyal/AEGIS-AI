"""
ocr_layer.py -- turns an attachment (image / PDF / DOCX / plain text) into TEXT. Runs FIRST.

Rules
  * The LLM never receives the raw image/file, only text that has gone through NeMo -> Laya -> vault.
  * File type is decided by magic bytes, not by the client's filename or MIME header.
  * Size / page / pixel caps stop decompression bombs and OCR denial-of-service.
  * If a file cannot be read, the caller must treat it as unreadable and BLOCK (we cannot inspect what we
    cannot read). Errors carry a reason code only, never file content.
Dependencies (lazy, only the ones you use): pytesseract + Pillow + the tesseract binary (add `hin` for Hindi),
pypdf (PDF text layer), pdf2image/poppler (scanned PDFs), python-docx.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 10
MAX_PIXELS = 25_000_000
MIN_TEXT_LAYER_CHARS = 30       # a PDF page with less text than this is treated as scanned


class OcrError(Exception):
    """Carries a reason code only (UNSUPPORTED_TYPE, TOO_LARGE, TOO_MANY_PAGES, ENGINE_FAILED, ...)."""
    def __init__(self, code: str):
        super().__init__(code); self.code = code


@dataclass
class Attachment:
    data: bytes
    filename: str = ""           # informational only; never trusted for type detection


@dataclass
class OcrResult:
    text: str
    kind: str                    # image | pdf | docx | text
    pages: int = 1
    confidence: float | None = None      # mean OCR word confidence 0-100 (None when no OCR was needed)
    used_ocr: bool = False
    warnings: list[str] = field(default_factory=list)


def sniff(data: bytes) -> str:
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff" or data[:4] in (b"GIF8", b"RIFF") \
            or data[:2] == b"BM" or data[:4] in (b"II*\x00", b"MM\x00*"):
        return "image"
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04" and b"word/" in data[:4096] + data[-4096:]:
        return "docx"
    try:
        data[:4096].decode("utf-8")
        if b"\x00" not in data[:4096]:
            return "text"
    except UnicodeDecodeError:
        pass
    return "unknown"


class TesseractEngine:
    """OCR for one PIL image. lang e.g. 'eng' or 'eng+hin' (needs the traineddata files installed)."""
    def __init__(self, lang: str = "eng", timeout_s: int = 30):
        self.lang, self.timeout_s = lang, timeout_s

    def image_to_text(self, img) -> tuple[str, float | None]:
        import pytesseract
        d = pytesseract.image_to_data(img, lang=self.lang, output_type=pytesseract.Output.DICT, timeout=self.timeout_s)
        words = [(w, float(c)) for w, c in zip(d["text"], d["conf"]) if w.strip() and float(c) >= 0]
        # rebuild lines so context windows stay readable
        lines: dict[tuple, list[str]] = {}
        for i, w in enumerate(d["text"]):
            if w.strip():
                lines.setdefault((d["block_num"][i], d["par_num"][i], d["line_num"][i]), []).append(w)
        text = "\n".join(" ".join(v) for v in lines.values())
        conf = sum(c for _, c in words) / len(words) if words else None
        return text, conf


class OcrLayer:
    def __init__(self, engine=None, max_bytes=MAX_BYTES, max_pages=MAX_PAGES, max_pixels=MAX_PIXELS):
        self.engine = engine or TesseractEngine()
        self.max_bytes, self.max_pages, self.max_pixels = max_bytes, max_pages, max_pixels

    # ---- public ---------------------------------------------------------- #
    def extract(self, att: Attachment) -> OcrResult:
        if len(att.data) > self.max_bytes:
            raise OcrError("TOO_LARGE")
        kind = sniff(att.data)
        try:
            if kind == "image":
                return self._image(att.data)
            if kind == "pdf":
                return self._pdf(att.data)
            if kind == "docx":
                return self._docx(att.data)
            if kind == "text":
                return OcrResult(att.data.decode("utf-8", errors="replace"), "text")
        except OcrError:
            raise
        except Exception:                                  # never leak library messages: they can echo content
            raise OcrError("ENGINE_FAILED") from None
        raise OcrError("UNSUPPORTED_TYPE")

    # ---- per type -------------------------------------------------------- #
    def _open_image(self, data: bytes):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = self.max_pixels
        img = Image.open(io.BytesIO(data))
        if img.width * img.height > self.max_pixels:
            raise OcrError("TOO_LARGE")
        return img

    def _image(self, data: bytes) -> OcrResult:
        img = self._open_image(data)
        n = getattr(img, "n_frames", 1)
        if n > self.max_pages:
            raise OcrError("TOO_MANY_PAGES")
        texts, confs = [], []
        for i in range(n):
            img.seek(i)
            t, c = self.engine.image_to_text(img.convert("RGB"))
            texts.append(t)
            if c is not None:
                confs.append(c)
        return OcrResult("\n".join(texts).strip(), "image", n, sum(confs) / len(confs) if confs else None, True)

    def _pdf(self, data: bytes) -> OcrResult:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise OcrError("ENCRYPTED_PDF")
        n = len(reader.pages)
        if n > self.max_pages:
            raise OcrError("TOO_MANY_PAGES")
        texts, confs, used_ocr, warn = [], [], False, []
        for i, page in enumerate(reader.pages):
            t = (page.extract_text() or "").strip()
            if len(t) < MIN_TEXT_LAYER_CHARS:               # scanned page: rasterize + OCR
                from pdf2image import convert_from_bytes
                for img in convert_from_bytes(data, dpi=200, first_page=i + 1, last_page=i + 1):
                    t, c = self.engine.image_to_text(img.convert("RGB"))
                    used_ocr = True
                    if c is not None:
                        confs.append(c)
            texts.append(t)
        return OcrResult("\n\n".join(texts).strip(), "pdf", n, sum(confs) / len(confs) if confs else None,
                         used_ocr, warn)

    def _docx(self, data: bytes) -> OcrResult:
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        used_ocr, confs = False, []
        for rel in d.part.rels.values():                    # images embedded in the document are OCR'd too
            if "image" in rel.reltype:
                r = self._image(rel.target_part.blob)
                parts.append(r.text); used_ocr = True
                if r.confidence is not None:
                    confs.append(r.confidence)
        warn = []
        return OcrResult(re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip(), "docx", 1,
                         sum(confs) / len(confs) if confs else None, used_ocr, warn)
